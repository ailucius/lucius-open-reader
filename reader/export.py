"""The requirement list as a bid checklist: an Excel workbook (Checklist + Summary) and a CSV.

Built on the server from a finished run's result, so the file holds exactly what the page shows:
every requirement with its verbatim quote and page, the decider's judgment when there is one, the
quote check (including "scanned page" when the quote was checked against the vision model's
transcription of a page image), and an empty Done column for the bidder.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import UTC, datetime
from urllib.parse import parse_qsl, unquote, urlsplit

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from .decide import deciding_from
from .models import Requirement, RunResult

COLUMNS: list[tuple[str, int]] = [  # (header, width in characters)
    ("No.", 6),
    ("Requirement", 48),
    ("In the tender's words", 60),
    ("Page", 7),
    ("Category", 13),
    ("Mandatory", 11),
    ("Judgment", 14),
    ("Why", 40),
    ("Quote check", 24),
    ("Done", 9),
]
JUDGMENT = {"fails": "fails", "not stated": "to confirm", "met": "met", "routine": "routine form"}
VERDICT = {"bid": "Bid", "bid with conditions": "Bid, with conditions", "no bid": "No bid"}
FIT = {"fits": "your line of work", "partly": "partly your line of work", "outside": "outside your line of work"}
NOT_A_SUBSTITUTE = (
    "This checklist is a reading aid made with AI models. It does not replace reading the tender: "
    "check every item against the tender documents before you bid."
)
# the page's own colours (light theme): text, background
COLOURS = {
    "fails": ("A8261E", "F7E1DF"),
    "to confirm": ("9A5600", "F8EBD6"),
    "met": ("1D6E45", "E1F0E7"),
    "not found": ("A8261E", "F7E1DF"),
}
HEADER = ("FFFFFF", "1F4F8B")
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def check_label(r: Requirement) -> str:
    if r.verified and r.joined:
        label = "found, joined"
    elif r.verified and r.cited_page:
        label = "found, page fixed"
    elif r.verified:
        label = "found"
    else:
        label = "not found"
    return label + ", scanned page" if r.scanned else label


def tender_name(result: RunResult) -> str:
    """The tender's title when the decider read one from its first pages, else the file's name."""
    if result.decision and result.decision.scope and result.decision.scope.title:
        return result.decision.scope.title
    parts = urlsplit(result.source)
    for key, value in parse_qsl(parts.query):
        if key.lower().endswith("filename") and value.strip():
            return re.sub(r"\.pdf$", "", value.strip(), flags=re.I)
    last = unquote(parts.path.rstrip("/").split("/")[-1]) if parts.path else result.source
    return re.sub(r"\.pdf$", "", last, flags=re.I) or "tender"


def export_filename(result: RunResult, ext: str) -> str:
    """ASCII only, so it needs no quoting in Content-Disposition."""
    ascii_name = unicodedata.normalize("NFKD", tender_name(result)).encode("ascii", "ignore").decode()
    slug = re.sub(r"[^a-z0-9]+", "-", ascii_name.lower()).strip("-")
    if len(slug) > 60:  # cut at a word, not inside one
        slug = slug[:61].rsplit("-", 1)[0] if "-" in slug[:61] else slug[:60]
    return f"lucius-checklist-{slug or 'tender'}.{ext}"


def checklist_rows(result: RunResult) -> list[list[str | int]]:
    """Header plus one row per requirement, in document order."""
    judged = {a.id: a for a in result.decision.assessments} if result.decision else {}
    rows: list[list[str | int]] = [[name for name, _ in COLUMNS]]
    for i, r in enumerate(result.requirements, start=1):
        a = judged.get(r.id)
        why = ""
        if a:
            why = ("Hard to obtain. " if a.hard and a.status in ("fails", "not stated") else "") + a.why
        rows.append(
            [
                int(r.id[1:]) if r.id[1:].isdigit() else i,
                r.requirement,
                r.quote,
                r.page,
                r.category.capitalize(),
                "yes" if r.mandatory else "no",
                JUDGMENT[a.status] if a else "",
                why,
                check_label(r),
                "",
            ]
        )
    return rows


def page_list(numbers: list[int]) -> str:
    """[1, 2, 3, 7] -> "1 to 3, 7"."""
    runs: list[list[int]] = []
    for n in sorted(set(numbers)):
        if runs and n == runs[-1][1] + 1:
            runs[-1][1] = n
        else:
            runs.append([n, n])
    return ", ".join(str(a) if a == b else f"{a}, {b}" if b == a + 1 else f"{a} to {b}" for a, b in runs)


def summary_rows(result: RunResult, generated: datetime) -> list[tuple[str, str]]:
    s, d = result.stats, result.decision
    scope = d.scope if d else None
    rows: list[tuple[str, str]] = [("Tender", tender_name(result))]
    if scope and scope.buyer:
        rows.append(("Buyer", scope.buyer))
    rows.append(("Source", result.source))
    if result.read_at:
        rows.append(("Read on", result.read_at.replace("T", " ").replace("+00:00", " UTC")))
    rows.append(("Pages", f"{s.pages_total} pages, {s.pages_read} read"))
    if s.scanned_pages:
        rows.append(
            (
                "Scanned pages",
                f"{len(s.scanned_pages)} read from the page image by {result.vision_model or 'the vision model'}: "
                f"pages {page_list(s.scanned_pages)}. These pages have no text layer, so their quotes were checked "
                "against the model's transcription of the page image, not against a text layer.",
            )
        )
    else:
        rows.append(("Scanned pages", "None"))
    if s.scan_failed_pages:
        rows.append(
            ("Scans not read", f"Pages {page_list(s.scan_failed_pages)}: the vision model could not read them in two tries.")
        )
    if s.scan_skipped_pages:
        why = f"the cap is {s.scan_cap} scanned pages a run" if s.scan_cap else "reading scanned pages is switched off"
        rows.append(("Scans over the cap", f"Pages {page_list(s.scan_skipped_pages)} are scans and were not read: {why}."))
    if s.empty_page_numbers:
        rows.append(("Blank pages", page_list(s.empty_page_numbers)))
    mandatory = sum(1 for r in result.requirements if r.mandatory)
    rows.append(
        (
            "Requirements",
            f"{s.requirements} ({mandatory} mandatory). Quotes found on their page: {s.verified} of {s.requirements}.",
        )
    )
    rows.append(("Company profile", result.profile.strip() or "None given"))
    if d:
        rows.append(("Call", VERDICT[d.verdict]))
        rows.append(("Confidence", f"{round(d.confidence * 100)}%"))
        if scope:
            rows.append(("Line of work", f"{FIT[scope.fit]}. {scope.buys}".strip()))
        c = d.counts
        rows.append(
            (
                "Judgments",
                f"{c.get('fails', 0)} fail, {c.get('not stated', 0)} to confirm, {c.get('met', 0)} met, "
                f"{c.get('routine', 0)} routine forms",
            )
        )
        if d.reasons:
            rows.append(("Why", "\n".join(d.reasons)))
        if d.conditions:
            rows.append(("Before you bid", "\n".join(f"{i}. {t}" for i, t in enumerate(d.conditions, start=1))))
        if result.decision_note:
            rows.append(("Note on the call", result.decision_note))
    else:
        rows.append(("Call", "No call. " + (result.decision_note or "")))
    models = [f"reader {result.reader_model}"]
    if d:
        models.append(f"decider {result.decider_model}")
    if result.vision_model:
        models.append(f"scanned pages {result.vision_model}")
    rows.append(("Models", ", ".join(models) + (" (offline fake client)" if result.fake_llm else "")))
    rows.append(("Generated", generated.strftime("%Y-%m-%d %H:%M UTC")))
    rows.append(("Note", NOT_A_SUBSTITUTE))
    return rows


DECIDING_COLUMNS: list[tuple[str, int]] = [  # under the summary: A and B keep the summary's widths
    ("No.", 24),
    ("Requirement", 90),
    ("Page", 7),
    ("Judgment", 14),
    ("Hard to obtain", 15),
    ("Why", 50),
]


def deciding_rows(result: RunResult) -> list[list[str | int]]:
    """Every requirement the company fails or must confirm: failures first, then the hardest to obtain."""
    if not result.decision:
        return []
    unmet = deciding_from(result.decision.assessments, result.requirements)
    return [
        [int(x.id[1:]) if x.id[1:].isdigit() else x.id, x.requirement, x.page, JUDGMENT[x.status], "yes" if x.hard else "no", x.why]
        for x in unmet
    ]


def _text(value: str | int) -> str | int:
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def _put(ws: Worksheet, row: int, col: int, value: str | int) -> Cell:
    cell = ws.cell(row=row, column=col, value=_text(value))
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"  # tender text, never a formula
    return cell


def _fill(cell: Cell, label: str) -> None:
    colours = COLOURS.get(label.split(",")[0])
    if colours:
        cell.font = Font(color=colours[0], bold=True)
        cell.fill = PatternFill("solid", fgColor=colours[1])


def _header(cell: Cell) -> None:
    cell.font = Font(color=HEADER[0], bold=True)
    cell.fill = PatternFill("solid", fgColor=HEADER[1])
    cell.alignment = Alignment(vertical="center", wrap_text=True)


def checklist_xlsx(result: RunResult, generated: datetime | None = None) -> bytes:
    generated = generated or datetime.now(UTC)
    wb = Workbook()
    ws = wb.active
    ws.title = "Checklist"
    rows = checklist_rows(result)
    top = Alignment(vertical="top", wrap_text=True)
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            cell = _put(ws, r, c, value)
            if r == 1:
                _header(cell)
            else:
                cell.alignment = top
                if COLUMNS[c - 1][0] in ("Judgment", "Quote check") and isinstance(value, str):
                    _fill(cell, value)
    for c, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.freeze_panes = "B2"
    last = get_column_letter(len(COLUMNS))
    ws.auto_filter.ref = f"A1:{last}{len(rows)}"
    if len(rows) > 1:
        done = DataValidation(type="list", formula1='"yes,no,not needed"', allow_blank=True)
        done.add(f"{last}2:{last}{len(rows)}")
        ws.add_data_validation(done)

    sm = wb.create_sheet("Summary")
    for c, (_, width) in enumerate(DECIDING_COLUMNS, start=1):
        sm.column_dimensions[get_column_letter(c)].width = width
    r = 1
    for label, value in summary_rows(result, generated):
        _put(sm, r, 1, label).font = Font(bold=True)
        _put(sm, r, 2, value).alignment = top
        sm.cell(row=r, column=1).alignment = Alignment(vertical="top")
        r += 1
    deciding = deciding_rows(result)
    if deciding:
        r += 1
        _put(sm, r, 1, "Requirements that decide it").font = Font(bold=True, size=12)
        r += 1
        for c, (head, _) in enumerate(DECIDING_COLUMNS, start=1):
            _header(_put(sm, r, c, head))
        for row in deciding:
            r += 1
            for c, value in enumerate(row, start=1):
                cell = _put(sm, r, c, value)
                cell.alignment = top
                if c == 4 and isinstance(value, str):
                    _fill(cell, value)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def checklist_csv(result: RunResult) -> bytes:
    """The Checklist sheet as CSV (UTF-8 with a byte order mark, so Excel reads accents right).
    A cell that starts like a formula gets a leading apostrophe, so a spreadsheet shows it as text."""
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\r\n")
    for row in checklist_rows(result):
        writer.writerow(
            ["'" + v if isinstance(v, str) and v.startswith(_FORMULA_START) else v for v in row]
        )
    return out.getvalue().encode("utf-8-sig")
