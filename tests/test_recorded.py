"""Replay of one real run with scanned pages, offline.

tests/fixtures/scan_run.json holds the replies Nebius Token Factory sent on 22 Oct 2026 when the
reader ran on a 4-page PDF whose pages 2 and 3 are scans (page 3 with a portal's page stamp on top):
Nemotron 3.5 Super VL transcribed the two page images, Nemotron 3 Super read the pages, Nemotron 3
Ultra made the call. The replay serves those replies through the real NebiusClient, so the whole
pipeline runs as it did live, with no network. Re-record with tests/fixtures/record_scan_run.py.
"""

from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any

import httpx
from openpyxl import load_workbook

from reader.config import DEFAULT_DECIDER_MODEL, DEFAULT_READER_MODEL, DEFAULT_VISION_MODEL, Settings
from reader.export import checklist_xlsx
from reader.llm import NebiusClient
from reader.pipeline import run_reader
from tests.fixtures.record_scan_run import PROFILE, recorded_pdf, replay_key

RECORDING = json.loads((Path(__file__).parent / "fixtures" / "scan_run.json").read_text())


class Replay(httpx.AsyncBaseTransport):
    """Serves each recorded reply to the request with the same replay key (and attempt number)."""

    def __init__(self) -> None:
        self.replies = {(c["key"], c["attempt"]): c for c in RECORDING["calls"]}
        self.attempts: dict[str, int] = {}
        self.bodies: list[dict[str, Any]] = []

    async def handle_async_request(self, request: httpx.Request) -> httpx.Response:
        assert request.url == RECORDING["base_url"].rstrip("/") + "/chat/completions"
        assert request.headers["authorization"] == "Bearer nb-test"
        body = json.loads(request.content)
        self.bodies.append(body)
        key = replay_key(body)
        self.attempts[key] = self.attempts.get(key, 0) + 1
        call = self.replies.get((key, self.attempts[key]))
        assert call is not None, f"no recorded reply for {key!r}: re-record tests/fixtures/scan_run.json"
        assert body["model"] == call["model"]
        return httpx.Response(call["status"], json=call["reply"])


async def _replay() -> tuple[Any, Replay, NebiusClient]:
    replay = Replay()
    client = NebiusClient("nb-test", RECORDING["base_url"], transport=replay)
    result, _ = await run_reader(
        settings=Settings(nebius_api_key="nb-test"), llm=client, profile=PROFILE, pdf=recorded_pdf(), source_name="W-2026-14.pdf"
    )
    return result, replay, client


def _reply_text(key: str) -> str:
    return next(c for c in RECORDING["calls"] if c["key"] == key)["reply"]["choices"][0]["message"]["content"]


async def test_replay_uses_every_recorded_call_once() -> None:
    _, replay, client = await _replay()
    assert sorted(replay.attempts.items()) == sorted(
        (k, max(c["attempt"] for c in RECORDING["calls"] if c["key"] == k)) for k in {c["key"] for c in RECORDING["calls"]}
    )
    assert client.usage == RECORDING["usage"]  # tokens counted from each reply's usage, per model
    assert set(client.usage) == {DEFAULT_VISION_MODEL, DEFAULT_READER_MODEL, DEFAULT_DECIDER_MODEL}


async def test_the_scanned_pages_go_to_the_vision_model_as_images() -> None:
    _, replay, _ = await _replay()
    images = [b for b in replay.bodies if isinstance(b["messages"][-1]["content"], list)]
    assert len(images) == 2 and all(b["model"] == DEFAULT_VISION_MODEL for b in images)
    for body in images:
        assert body["chat_template_kwargs"] == {"enable_thinking": False} and "response_format" not in body
        text, image = body["messages"][-1]["content"]
        assert text["type"] == "text" and image["image_url"]["url"].startswith("data:image/png;base64,")
    assert sorted(b["messages"][-1]["content"][0]["text"] for b in images) == [
        "Page 2 of the tender. Transcribe it.",
        "Page 3 of the tender. Transcribe it.",
    ]


async def test_the_reader_reads_the_vision_models_transcription() -> None:
    _, replay, _ = await _replay()
    reader = next(b for b in replay.bodies if b["model"] == DEFAULT_READER_MODEL)
    prompt = reader["messages"][-1]["content"]
    for page in (2, 3):
        transcription = _reply_text(f"image | Page {page} of the tender. Transcribe it.").strip()
        assert f"=== PAGE {page} ===\n{transcription}\n" in prompt
    assert "Page 4 of 9" in prompt  # the portal's stamp on page 3 was transcribed with the page


async def test_quotes_from_scanned_pages_are_checked_against_the_transcription_and_marked() -> None:
    result, _, _ = await _replay()
    recorded = RECORDING["result"]
    assert result.stats.model_dump() == recorded["stats"]
    assert result.stats.scanned_pages == [2, 3] and result.stats.pages_read == 4
    assert result.vision_model == DEFAULT_VISION_MODEL
    assert [r.model_dump() for r in result.requirements] == recorded["requirements"]
    by_page = {p: [r for r in result.requirements if r.page == p] for p in (1, 2, 3, 4)}
    assert by_page[2] and by_page[3] and by_page[4]
    assert all(r.scanned and r.verified for p in (2, 3) for r in by_page[p])
    assert not any(r.scanned for p in (1, 4) for r in by_page[p])
    assert result.decision is not None and result.decision.verdict == recorded["decision"]["verdict"]

    ws = load_workbook(io.BytesIO(checklist_xlsx(result)))["Checklist"]
    checks = {(row[3], row[8]) for row in ws.iter_rows(min_row=2, values_only=True)}
    assert {c for p, c in checks if p in (2, 3)} == {"found, scanned page"}
    assert {c for p, c in checks if p == 4} == {"found"}
