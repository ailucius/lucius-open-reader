"""The bid checklist export: the workbook and CSV are opened and read back, and the download routes."""

from __future__ import annotations

import csv
import io
import json
from datetime import UTC, datetime

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from reader.app import EXAMPLE, create_app
from reader.config import Settings
from reader.export import COLUMNS, checklist_csv, checklist_rows, checklist_xlsx, export_filename, page_list
from reader.llm import FakeLLM
from reader.models import Assessment, Decision, ReadStats, Requirement, RunResult, Scope
from tests.conftest import TENDER_PAGES, make_scanned_pdf

GENERATED = datetime(2026, 10, 21, 18, 30, tzinfo=UTC)


def _result() -> RunResult:
    reqs = [
        Requirement(id="R1", requirement="Register at grade 7.", quote="Bidders must be registered at grade 7 or higher.", page=2,
                    category="eligibility", verified=True, match_score=1.0, scanned=True),
        Requirement(id="R2", requirement="Pay a bid guarantee.", quote="=A bid guarantee of 2% must be submitted.", page=4,
                    category="commercial", verified=False, match_score=0.4),
        Requirement(id="R3", requirement="Sign Annex C.", quote="Bidders must sign\x0b the declaration (Annex C).", page=6,
                    category="legal", verified=True, match_score=1.0, cited_page=5),
        Requirement(id="R4", requirement="Variable speed drives.", quote="should preferably use variable speed drives", page=3,
                    category="technical", mandatory=False, verified=True, match_score=1.0),
    ]
    decision = Decision(
        verdict="no bid",
        confidence=0.8,
        scope=Scope(title="Bulk water supply pipeline", buyer="City of Example Water Department", buys="A pipeline.", fit="fits", why="w"),
        assessments=[
            Assessment(id="R1", status="not stated", hard=True, why="Grade not in the profile."),
            Assessment(id="R2", status="fails", hard=True, why="No bank facility."),
            Assessment(id="R3", status="routine", why=""),
        ],
        counts={"fails": 1, "not stated": 1, "met": 0, "routine": 1},
        reasons=["The company cannot raise the guarantee [R2]."],
        model="decider",
    )
    stats = ReadStats(
        pages_total=9, pages_read=7, requirements=4, verified=3, unverified=1,
        scanned_pages=[2, 3], scan_failed_pages=[8], scan_skipped_pages=[9], scan_cap=2, empty_page_numbers=[7],
    )
    return RunResult(
        source="https://buyer.example/docs/W-14%20tender.pdf", pdf_url="https://buyer.example/docs/W-14%20tender.pdf",
        profile="Pipeline contractor, grade 6.", reader_model="reader", decider_model="decider",
        vision_model="nvidia/Nemotron-3.5-Super-VL", fake_llm=False, read_at="2026-10-21T18:02:11+00:00",
        stats=stats, requirements=reqs, decision=decision,
    )


def test_checklist_sheet_cells() -> None:
    wb = load_workbook(io.BytesIO(checklist_xlsx(_result(), GENERATED)))
    assert wb.sheetnames == ["Checklist", "Summary"]
    ws = wb["Checklist"]
    assert [c.value for c in ws[1]] == [name for name, _ in COLUMNS]
    assert ws.max_row == 5 and ws.freeze_panes == "B2" and ws.auto_filter.ref == "A1:J5"
    row = {name: ws.cell(row=2, column=i).value for i, (name, _) in enumerate(COLUMNS, start=1)}
    assert row["No."] == 1 and row["Page"] == 2 and row["Category"] == "Eligibility" and row["Mandatory"] == "yes"
    assert row["In the tender's words"] == "Bidders must be registered at grade 7 or higher."
    assert row["Judgment"] == "to confirm" and row["Why"] == "Hard to obtain. Grade not in the profile."
    assert row["Quote check"] == "found, scanned page" and row["Done"] is None
    # tender text that starts like a formula stays text; control characters PDFs carry are dropped
    quote = ws.cell(row=3, column=3)
    assert quote.value == "=A bid guarantee of 2% must be submitted." and quote.data_type == "s"
    assert [ws.cell(row=3, column=c).value for c in (7, 9)] == ["fails", "not found"]
    assert ws.cell(row=4, column=3).value == "Bidders must sign the declaration (Annex C)."
    assert [ws.cell(row=4, column=c).value for c in (7, 9)] == ["routine form", "found, page fixed"]
    assert [ws.cell(row=5, column=c).value for c in (6, 7, 8)] == ["no", None, None]  # optional: not judged
    assert str(ws.data_validations.dataValidation[0].sqref) == "J2:J5"  # the Done column offers yes, no, not needed


def test_summary_sheet_says_what_was_read_and_what_decides() -> None:
    wb = load_workbook(io.BytesIO(checklist_xlsx(_result(), GENERATED)))
    rows = [tuple(r) for r in wb["Summary"].iter_rows(values_only=True)]
    summary = {r[0]: r[1] for r in rows if r[0] and r[1]}
    assert summary["Tender"] == "Bulk water supply pipeline" and summary["Buyer"] == "City of Example Water Department"
    assert summary["Read on"] == "2026-10-21 18:02:11 UTC" and summary["Pages"] == "9 pages, 7 read"
    assert "2 read from the page image by nvidia/Nemotron-3.5-Super-VL: pages 2, 3" in summary["Scanned pages"]
    assert "transcription of the page image, not against a text layer" in summary["Scanned pages"]
    assert summary["Scans not read"].startswith("Pages 8:")
    assert summary["Scans over the cap"] == "Pages 9 are scans and were not read: the cap is 2 scanned pages a run."
    assert summary["Blank pages"] == "7"
    assert summary["Call"] == "No bid" and summary["Confidence"] == "80%"
    assert summary["Judgments"] == "1 fail, 1 to confirm, 0 met, 1 routine forms"
    assert summary["Generated"] == "2026-10-21 18:30 UTC"
    assert "does not replace reading the tender" in summary["Note"]
    start = next(i for i, r in enumerate(rows) if r[0] == "Requirements that decide it")
    assert rows[start + 1] == ("No.", "Requirement", "Page", "Judgment", "Hard to obtain", "Why")
    assert rows[start + 2][:4] == (2, "Pay a bid guarantee.", 4, "fails")  # failures first
    assert rows[start + 3][:4] == (1, "Register at grade 7.", 2, "to confirm")
    assert len(rows) == start + 4


def test_checklist_without_a_decision_leaves_judgments_empty() -> None:
    result = _result()
    result.decision = None
    result.decision_note = "Add a company profile to get a bid / no-bid call."
    rows = checklist_rows(result)
    assert all(r[6] == "" and r[7] == "" for r in rows[1:])
    wb = load_workbook(io.BytesIO(checklist_xlsx(result, GENERATED)))
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(values_only=True) if r[0]}
    assert summary["Call"] == "No call. Add a company profile to get a bid / no-bid call."
    assert "Requirements that decide it" not in summary


def test_csv_is_the_checklist_sheet() -> None:
    raw = checklist_csv(_result())
    assert raw.startswith(b"\xef\xbb\xbf")  # so Excel reads it as UTF-8
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    assert rows[0] == [name for name, _ in COLUMNS] and len(rows) == 5
    assert rows[1][:4] == ["1", "Register at grade 7.", "Bidders must be registered at grade 7 or higher.", "2"]
    assert rows[1][8] == "found, scanned page"
    assert rows[2][2] == "'=A bid guarantee of 2% must be submitted."  # shown as text, not run as a formula


def test_filenames_and_page_lists() -> None:
    result = _result()
    assert export_filename(result, "xlsx") == "lucius-checklist-bulk-water-supply-pipeline.xlsx"
    result.decision = None
    assert export_filename(result, "csv") == "lucius-checklist-w-14-tender.csv"
    result.source = "https://www.etenders.gov.za/home/Download?blobName=x.pdf&downloadedFileName=TENDER%20DOCUMENT%20-%20WATER%20PUMPS%208-2-1-656.pdf"
    assert export_filename(result, "xlsx") == "lucius-checklist-tender-document-water-pumps-8-2-1-656.xlsx"
    result.source = "Évaluation des offres " + "très longue " * 10 + ".pdf"
    name = export_filename(result, "xlsx")
    assert name.startswith("lucius-checklist-evaluation-des-offres-tres-longue") and len(name) <= 17 + 60 + 5
    assert not name[:-5].endswith("-") and "-tr." not in name
    assert page_list([9, 1, 2, 3, 7, 8]) == "1 to 3, 7 to 9" and page_list([4, 5]) == "4, 5"


def _events(text: str) -> list[dict[str, object]]:
    return [json.loads(line[6:]) for line in text.splitlines() if line.startswith("data: ")]


def test_download_routes_after_a_run_with_scanned_pages() -> None:
    pdf = make_scanned_pdf([("text", TENDER_PAGES[0]), ("scan", TENDER_PAGES[1]), ("scan", TENDER_PAGES[3])], dpi=60)
    app = create_app(Settings(fake_llm=True, chunk_pages=3, scan_dpi=60), llm=FakeLLM({2: TENDER_PAGES[1], 3: TENDER_PAGES[3]}))
    with TestClient(app) as client:
        run_id = client.post("/api/runs", data={"profile": "Pipeline contractor"}, files={"file": ("W-14.pdf", pdf, "application/pdf")}).json()["run_id"]
        with client.stream("GET", f"/api/runs/{run_id}/events") as resp:
            events = _events("".join(resp.iter_text()))
        assert events[-1]["type"] == "result"
        assert [e["page"] for e in events if e["type"] == "scan"] in ([2, 3], [3, 2])

        xlsx = client.get(f"/api/runs/{run_id}/checklist.xlsx")
        assert xlsx.status_code == 200
        assert xlsx.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert xlsx.headers["content-disposition"] == 'attachment; filename="lucius-checklist-w-14.xlsx"'
        ws = load_workbook(io.BytesIO(xlsx.content))["Checklist"]
        checks = {(row[3], row[8]) for row in ws.iter_rows(min_row=2, values_only=True)}
        assert ("found, scanned page" in {c for p, c in checks if p in (2, 3)}) and all(c == "found" for p, c in checks if p == 1)
        assert all(row[6] for row in ws.iter_rows(min_row=2, values_only=True) if row[5] == "yes")  # every mandatory item judged

        got = client.get(f"/api/runs/{run_id}/checklist.csv")
        assert got.status_code == 200 and got.headers["content-type"].startswith("text/csv")
        assert got.headers["content-disposition"] == 'attachment; filename="lucius-checklist-w-14.csv"'
        assert got.content.decode("utf-8-sig").splitlines()[0].startswith("No.,Requirement,")

        assert client.get(f"/api/runs/{run_id}/checklist.pdf").status_code == 404
        assert client.get("/api/runs/nope/checklist.xlsx").status_code == 404


def test_example_checklist_route() -> None:
    with TestClient(create_app(Settings(fake_llm=True))) as client:
        resp = client.get("/api/example/checklist.xlsx")
        if not EXAMPLE.is_file():
            assert resp.status_code == 404
            return
        assert resp.status_code == 200 and "attachment" in resp.headers["content-disposition"]
        ws = load_workbook(io.BytesIO(resp.content))["Checklist"]
        saved = json.loads(EXAMPLE.read_text())
        assert ws.max_row == len(saved["requirements"]) + 1
